import os
import sys
import json
import re

if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from python_calamine import CalamineWorkbook
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
INDEX_PATH = os.path.join(r"C:\Users\DELL\.gemini\antigravity\brain\f0678079-bf49-4bdd-be79-d5e6298486ad\scratch", "tt38_index.json")

if not os.path.exists(INDEX_PATH):
    INDEX_PATH = os.path.join(SCRIPT_DIR, "tt38_index.json")

with open(INDEX_PATH, "r", encoding="utf-8") as f:
    TT38_INDEX = json.load(f)

# Bản đồ tiền tố
PREFIX_MAP = {}
for code, info in TT38_INDEX.items():
    clean = code.strip().upper()
    PREFIX_MAP[clean] = info
    parts = clean.split('.')
    if len(parts) == 2:
        prefix_alpha, num_part = parts[0], parts[1]
        for l in range(3, len(num_part)):
            sub = f"{prefix_alpha}.{num_part[:l]}"
            if sub not in PREFIX_MAP:
                PREFIX_MAP[sub] = info

def find_tt38_reference(raw_code):
    if not raw_code:
        return None
    code = str(raw_code).strip().upper()
    
    # Bỏ qua các ký tự rác hoặc tiêu đề mục
    if code in ['', '*', '-', 'STT', 'MÃ SỐ', 'MÃ HIỆU', 'NONE']:
        return None
        
    # Mã tạm tính
    if code in ['TT', 'TẠM TÍNH', 'TAMTINH', 'TẠMTÍNH']:
        return {
            'code': code,
            'match_type': 'Tạm tính (Mã TT)',
            'status': 'Tạm tính',
            'appendix': None,
            'pdf_file': None,
            'pdf_path': None,
            'page': None
        }

    # 1. Khớp chính xác
    if code in TT38_INDEX:
        info = TT38_INDEX[code]
        pages = [p for p in info['pages'] if p < 700] or info['pages']
        return {
            'code': code,
            'match_type': 'Khớp chính xác',
            'status': 'Đã khớp TT38',
            'appendix': info['appendix'],
            'pdf_file': info['pdf_file'],
            'pdf_path': info['pdf_path'],
            'page': pages[0]
        }
        
    # 2. Khớp tiền tố
    if '.' in code:
        alpha, digits = code.split('.', 1)
        for length in range(len(digits) - 1, 2, -1):
            sub_code = f"{alpha}.{digits[:length]}"
            if sub_code in PREFIX_MAP:
                info = PREFIX_MAP[sub_code]
                pages = [p for p in info['pages'] if p < 700] or info['pages']
                return {
                    'code': sub_code,
                    'match_type': f'Khớp nhóm {sub_code}',
                    'status': 'Đã khớp TT38',
                    'appendix': info['appendix'],
                    'pdf_file': info['pdf_file'],
                    'pdf_path': info['pdf_path'],
                    'page': pages[0]
                }
            zero_code = f"{alpha}.{digits[:length]}{'0'*(5-length)}"
            if zero_code in TT38_INDEX:
                info = TT38_INDEX[zero_code]
                pages = [p for p in info['pages'] if p < 700] or info['pages']
                return {
                    'code': zero_code,
                    'match_type': f'Khớp nhóm {zero_code}',
                    'status': 'Đã khớp TT38',
                    'appendix': info['appendix'],
                    'pdf_file': info['pdf_file'],
                    'pdf_path': info['pdf_path'],
                    'page': pages[0]
                }

    # Nếu có dạng mã định mức chữ cái + số nhưng không thấy trong TT38
    if re.match(r'^[A-Z]{2}\.\d+', code):
        return {
            'code': code,
            'match_type': 'Chưa thấy trong TT38',
            'status': 'Chưa khớp',
            'appendix': None,
            'pdf_file': None,
            'pdf_path': None,
            'page': None
        }

    return None

def get_tt38_url(ref):
    if not ref or not ref.get('page'):
        return None
    pl_key = "pl2"
    app_str = str(ref.get('appendix', ''))
    if "VIII" in app_str: pl_key = "pl8"
    elif "VII" in app_str: pl_key = "pl7"
    elif "VI" in app_str: pl_key = "pl6"
    elif "V" in app_str: pl_key = "pl5"
    elif "IV" in app_str: pl_key = "pl4"
    elif "III" in app_str: pl_key = "pl3"
    elif "II" in app_str: pl_key = "pl2"
    elif "I" in app_str: pl_key = "pl1"
    return f"tt38://{pl_key}/{ref['page']}"

def detect_code_and_name_columns(rows):
    """Tự động phát hiện cột mã hiệu và cột tên công việc dựa vào nội dung và tiêu đề"""
    best_code_col = -1
    best_name_col = -1
    header_row_idx = -1
    
    # 1. Tìm theo tiêu đề trước (trong 15 dòng đầu)
    for r_idx, row in enumerate(rows[:15]):
        for c_idx, cell in enumerate(row):
            val_clean = re.sub(r'\s+', ' ', str(cell or '').strip().upper())
            if any(k in val_clean for k in ['MÃ HIỆU ĐỊNH MỨC', 'MÃ HIỆU CÔNG TÁC', 'MÃ HIỆU', 'MÃ SỐ', 'MÃ ĐƠN GIÁ']):
                if best_code_col == -1:
                    best_code_col = c_idx
                    header_row_idx = r_idx
            elif any(k in val_clean for k in ['TÊN CÔNG TÁC', 'DANH MỤC CÔNG TÁC', 'TÊN CÔNG VIỆC', 'NỘI DUNG CÔNG VIỆC', 'TÊN THIẾT BỊ']):
                if best_name_col == -1:
                    best_name_col = c_idx

    # 2. Nếu chưa tìm được bằng tiêu đề, quét tần suất xuất hiện mẫu mã hiệu (AA.xxx, TT...)
    if best_code_col == -1:
        max_matches = 0
        for c_idx in range(min(10, len(rows[0]) if rows else 0)):
            matches = 0
            for r in rows[:40]:
                if c_idx < len(r):
                    val = str(r[c_idx] or '').strip().upper()
                    if re.match(r'^[A-Z]{2}\.\d+', val) or val == 'TT':
                        matches += 1
            if matches > max_matches and matches >= 2:
                max_matches = matches
                best_code_col = c_idx
                if header_row_idx == -1:
                    header_row_idx = 3 # mặc định dòng 4

    if best_name_col == -1 and best_code_col != -1:
        # Cột tên thường nằm ngay sau cột mã số
        best_name_col = best_code_col + 1

    return best_code_col, best_name_col, header_row_idx

def link_estimate_file(input_path, output_path=None):
    if not output_path:
        base, ext = os.path.splitext(input_path)
        output_path = f"{base}_TT38_Linked.xlsx"

    print(f"\n--> ĐANG XỬ LÝ: {os.path.basename(input_path)}")
    wb_cal = CalamineWorkbook.from_path(input_path)
    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active) # Xóa sheet mặc định

    # Fonts & Styles
    link_font = Font(name='Times New Roman', size=11, color='0000FF', underline='single', bold=True)
    code_linked_font = Font(name='Times New Roman', size=11, color='002060', underline='single', bold=True)
    regular_font = Font(name='Times New Roman', size=11)
    bold_font = Font(name='Times New Roman', size=11, bold=True)
    
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid') # Xanh đậm sang trọng
    header_font_white = Font(name='Times New Roman', size=11, bold=True, color='FFFFFF')
    
    col_header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid') # Xanh nhạt
    tt_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid') # Vàng cảnh báo
    exact_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid') # Xanh lá nhạt khớp chuẩn
    
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    summary_items = []
    preferred_active_sheet = None

    for sheet_name in wb_cal.sheet_names:
        rows = wb_cal.get_sheet_by_name(sheet_name).to_python()
        if not rows:
            continue
            
        ws = out_wb.create_sheet(title=sheet_name[:31])
        code_col_idx, name_col_idx, header_row_idx = detect_code_and_name_columns(rows)
        
        has_items = False

        # Vị trí chèn cột link TT38: Chèn ngay sau cột Tên công tác (hoặc cột Mã số)
        # Để người dùng nhìn thấy NGAY LẬP TỨC trên màn hình mà không cần cuộn ngang!
        insert_link_after_col = name_col_idx if name_col_idx != -1 else (code_col_idx if code_col_idx != -1 else -1)

        for r_idx, row in enumerate(rows):
            excel_row = r_idx + 1
            col_offset = 0
            
            for c_idx, val in enumerate(row):
                excel_col = c_idx + 1 + col_offset
                cell_obj = ws.cell(row=excel_row, column=excel_col, value=val)
                cell_obj.font = regular_font
                
                # Nếu đây là ô Mã hiệu và có tham chiếu TT38 -> GẮN LINK TRỰC TIẾP LÊN MÃ HIỆU
                if code_col_idx != -1 and c_idx == code_col_idx and r_idx > header_row_idx:
                    ref = find_tt38_reference(val)
                    if ref and ref['page']:
                        file_url = get_tt38_url(ref)
                        cell_obj.hyperlink = file_url
                        cell_obj.font = code_linked_font
                        cell_obj.fill = exact_fill
                    elif ref and ref['status'] == 'Tạm tính':
                        cell_obj.fill = tt_fill

                # Nếu vừa ghi xong cột insert_link_after_col, chèn cột LINK TT38 ngay cạnh
                if insert_link_after_col != -1 and c_idx == insert_link_after_col:
                    link_col = excel_col + 1
                    col_offset = 1 # Dịch tất cả các cột phía sau sang phải 1 cột
                    
                    link_cell = ws.cell(row=excel_row, column=link_col)
                    
                    if r_idx == header_row_idx:
                        link_cell.value = "LINK TRA CỨU TT 38/2026/TT-BXD"
                        link_cell.font = bold_font
                        link_cell.fill = col_header_fill
                        link_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    elif r_idx > header_row_idx:
                        code_val = row[code_col_idx]
                        name_val = row[name_col_idx] if (name_col_idx != -1 and name_col_idx < len(row)) else ''
                        ref = find_tt38_reference(code_val)
                        
                        if ref:
                            has_items = True
                            if ref['page']:
                                file_url = get_tt38_url(ref)
                                display_text = f"👉 {ref['appendix']} - Trang {ref['page']} ({ref['match_type']})"
                                link_cell.value = f'=HYPERLINK("{file_url}", "{display_text}")'
                                link_cell.hyperlink = file_url
                                link_cell.font = link_font
                                link_cell.fill = exact_fill
                                
                                # Thêm vào danh sách dashboard
                                summary_items.append({
                                    'sheet': sheet_name,
                                    'row': excel_row,
                                    'code': str(code_val).strip(),
                                    'name': str(name_val).strip()[:70],
                                    'status': ref['status'],
                                    'appendix': ref['appendix'],
                                    'page': ref['page'],
                                    'match_type': ref['match_type'],
                                    'url': file_url
                                })
                            elif ref['status'] == 'Tạm tính':
                                link_cell.value = "⚠️ Tạm tính (Chưa có mã TT38)"
                                link_cell.font = regular_font
                                link_cell.fill = tt_fill
                                summary_items.append({
                                    'sheet': sheet_name,
                                    'row': excel_row,
                                    'code': str(code_val).strip(),
                                    'name': str(name_val).strip()[:70],
                                    'status': 'Tạm tính',
                                    'appendix': '-',
                                    'page': '-',
                                    'match_type': 'Tạm tính (Mã TT)',
                                    'url': None
                                })
                            link_cell.alignment = Alignment(horizontal='left', vertical='center')

        # Auto width
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col[:35]:
                if cell.value:
                    val_str = str(cell.value)
                    if not val_str.startswith('='):
                        max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 50), 12)

        if insert_link_after_col != -1:
            ws.column_dimensions[get_column_letter(insert_link_after_col + 2)].width = 38

        # Chọn sheet ưu tiên làm Active Sheet
        if has_items and not preferred_active_sheet:
            if any(k in sheet_name.upper() for k in ['CÔNG TRÌNH', 'GIÁ TỔNG HỢP', 'ĐƠN GIÁ CHI TIẾT']):
                preferred_active_sheet = ws

    # TẠO SHEET ĐẦU TIÊN: "BẢNG TRA CỨU TỔNG HỢP TT38"
    summary_ws = out_wb.create_sheet(title="TRA CỨU TT38 TỔNG HỢP", index=0)
    
    # Tiêu đề lớn
    summary_ws.merge_cells('A2:H2')
    title_cell = summary_ws.cell(row=2, column=1, value="BẢNG TỔNG HỢP LIÊN KẾT TRA CỨU ĐỊNH MỨC THÔNG TƯ 38/2026/TT-BXD")
    title_cell.font = Font(name='Times New Roman', size=16, bold=True, color='1F4E79')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    
    summary_ws.merge_cells('A3:H3')
    sub_cell = summary_ws.cell(row=3, column=1, value=f"Tệp dự toán: {os.path.basename(input_path)}  |  Tổng số hạng mục nhận diện: {len(summary_items)}  |  Bấm trực tiếp vào đường link để mở trang PDF Thông tư 38")
    sub_cell.font = Font(name='Times New Roman', size=11, italic=True)
    sub_cell.alignment = Alignment(horizontal='center', vertical='center')

    headers = ['STT', 'Sheet nguồn', 'Dòng', 'Mã hiệu dự toán', 'Tên công tác / Hạng mục', 'Trạng thái', 'Phân loại khớp', 'LINK MỞ TRANG THÔNG TƯ 38']
    for c_idx, h_text in enumerate(headers):
        h_cell = summary_ws.cell(row=5, column=c_idx + 1, value=h_text)
        h_cell.font = header_font_white
        h_cell.fill = header_fill
        h_cell.alignment = Alignment(horizontal='center', vertical='center')
        
    for idx, item in enumerate(summary_items):
        r = idx + 6
        summary_ws.cell(row=r, column=1, value=idx + 1).alignment = Alignment(horizontal='center')
        summary_ws.cell(row=r, column=2, value=item['sheet'])
        summary_ws.cell(row=r, column=3, value=item['row']).alignment = Alignment(horizontal='center')
        
        # Mã hiệu
        c_code = summary_ws.cell(row=r, column=4, value=item['code'])
        c_code.font = bold_font
        c_code.alignment = Alignment(horizontal='center')
        
        summary_ws.cell(row=r, column=5, value=item['name'])
        
        c_stat = summary_ws.cell(row=r, column=6, value=item['status'])
        c_stat.alignment = Alignment(horizontal='center')
        if item['status'] == 'Đã khớp TT38':
            c_stat.fill = exact_fill
        else:
            c_stat.fill = tt_fill
            
        summary_ws.cell(row=r, column=7, value=item['match_type'])
        
        # Link mở trang
        c_link = summary_ws.cell(row=r, column=8)
        if item['url']:
            link_label = f"👉 Mở {item['appendix']} - Trang {item['page']}"
            c_link.value = f'=HYPERLINK("{item['url']}", "{link_label}")'
            c_link.hyperlink = item['url']
            c_link.font = link_font
        else:
            c_link.value = "⚠️ Mã tạm tính - Kiểm tra báo giá"
            c_link.font = regular_font
            c_link.fill = tt_fill

    # Đặt độ rộng cột cho sheet tổng hợp
    summary_ws.column_dimensions['A'].width = 8
    summary_ws.column_dimensions['B'].width = 20
    summary_ws.column_dimensions['C'].width = 10
    summary_ws.column_dimensions['D'].width = 18
    summary_ws.column_dimensions['E'].width = 50
    summary_ws.column_dimensions['F'].width = 18
    summary_ws.column_dimensions['G'].width = 24
    summary_ws.column_dimensions['H'].width = 40

    # Luôn đặt Sheet đầu tiên (TRA CỨU TT38 TỔNG HỢP) làm Active Sheet!
    out_wb.active = summary_ws

    out_wb.save(output_path)
    matched_count = sum(1 for x in summary_items if x['url'])
    print(f"[THÀNH CÔNG] Đã lưu tệp: {os.path.basename(output_path)}")
    print(f"             -> Tổng công tác: {len(summary_items)} | Gắn link thành công: {matched_count} ({matched_count*100//max(len(summary_items),1)}%)")
    return output_path

if __name__ == "__main__":
    if len(sys.argv) > 1:
        for arg in sys.argv[1:]:
            if os.path.isfile(arg):
                link_estimate_file(arg)
    else:
        sample_folder = os.path.join(SCRIPT_DIR, "05.4.2_L1")
        if os.path.exists(sample_folder):
            for f in os.listdir(sample_folder):
                if (f.endswith('.xls') or f.endswith('.xlsx')) and not f.endswith('_Linked.xlsx'):
                    p = os.path.join(sample_folder, f)
                    try:
                        link_estimate_file(p)
                    except Exception as e:
                        print(f"[Lỗi] Không xử lý được {f}: {e}")
