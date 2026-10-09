# -*- coding: utf-8 -*-
"""
AUDIT ENGINE - Module phân tích, thẩm tra dự toán, đối chiếu Định mức TT38 và Báo giá
Tự động tính toán giảm trừ chi phí, gắn link TT38 và Báo giá, tô màu trực quan
"""

import os
import re
import json
import unicodedata
from python_calamine import CalamineWorkbook
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Style định dạng Excel thẩm tra
FILL_HEADER = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
FILL_SUBHEADER = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
FILL_GREEN = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Hợp lệ
FILL_YELLOW = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Tạm tính TT
FILL_RED = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")    # Vượt giá/Giảm trừ

FONT_TITLE = Font(name="Times New Roman", size=14, bold=True, color="1F4E79")
FONT_HEADER = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
FONT_REGULAR = Font(name="Times New Roman", size=10)
FONT_BOLD = Font(name="Times New Roman", size=10, bold=True)
FONT_LINK = Font(name="Times New Roman", size=10, color="0000FF", underline="single")
FONT_ALERT = Font(name="Times New Roman", size=10, bold=True, color="C00000")
FONT_OK = Font(name="Times New Roman", size=10, color="375623")
FONT_WARN = Font(name="Times New Roman", size=10, color="B25900")

BORDER_THIN = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

def get_pl_number(appendix_str):
    """Trích xuất số thứ tự phụ lục từ chuỗi La Mã (Phụ lục I..VIII) hoặc số"""
    if not appendix_str:
        return '2'
    m_roman = re.search(r'\b(VIII|VII|VI|V|IV|III|II|I)\b', str(appendix_str).upper())
    if m_roman:
        roman_map = {'I': '1', 'II': '2', 'III': '3', 'IV': '4', 'V': '5', 'VI': '6', 'VII': '7', 'VIII': '8'}
        return roman_map.get(m_roman.group(1), '2')
    m_num = re.search(r'\d+', str(appendix_str))
    if m_num:
        return m_num.group(0)
    return '2'

class AuditEngine:
    def __init__(self, tt38_index_path, quotation_indexer):
        self.tt38_index_path = tt38_index_path
        self.quotation_indexer = quotation_indexer
        self.tt38_index = self._load_tt38_index()
        self.prefix_map = self._build_prefix_map()

    def _load_tt38_index(self):
        with open(self.tt38_index_path, 'r', encoding='utf-8') as f:
            return json.load(f)

    def _build_prefix_map(self):
        p_map = {}
        for code, info in self.tt38_index.items():
            clean = code.strip().upper()
            p_map[clean] = info
            parts = clean.split('.')
            if len(parts) == 2:
                prefix_alpha, num_part = parts[0], parts[1]
                for l in range(3, len(num_part)):
                    sub = f"{prefix_alpha}.{num_part[:l]}"
                    if sub not in p_map:
                        p_map[sub] = info
        return p_map

    def find_tt38_norm(self, raw_code):
        """Tra cứu mã định mức trong TT38"""
        if not raw_code:
            return None
        code = str(raw_code).strip().upper()
        if code in ['', '*', '-', 'STT', 'MÃ SỐ', 'MÃ HIỆU', 'NONE', 'HM']:
            return None

        if code in ['TT', 'TẠM TÍNH', 'TAMTINH', 'TẠMTÍNH']:
            return {
                'code': code,
                'match_type': 'Tạm tính (Mã TT)',
                'status': 'Tạm tính',
                'appendix': None,
                'pdf_file': None,
                'page': None
            }

        # Khớp chính xác
        if code in self.tt38_index:
            info = self.tt38_index[code]
            pages = [p for p in info['pages'] if p < 700] or info['pages']
            return {
                'code': code,
                'match_type': 'Khớp chính xác',
                'status': 'Đã khớp TT38',
                'appendix': info['appendix'],
                'pdf_file': info['pdf_file'],
                'page': pages[0]
            }

        # Khớp tiền tố
        if '.' in code:
            alpha, digits = code.split('.', 1)
            for length in range(len(digits) - 1, 2, -1):
                sub_code = f"{alpha}.{digits[:length]}"
                if sub_code in self.prefix_map:
                    info = self.prefix_map[sub_code]
                    pages = [p for p in info['pages'] if p < 700] or info['pages']
                    return {
                        'code': sub_code,
                        'match_type': f'Khớp nhóm {sub_code}',
                        'status': 'Đã khớp TT38',
                        'appendix': info['appendix'],
                        'pdf_file': info['pdf_file'],
                        'page': pages[0]
                    }
                zero_code = f"{alpha}.{digits[:length]}{'0'*(5-length)}"
                if zero_code in self.tt38_index:
                    info = self.tt38_index[zero_code]
                    pages = [p for p in info['pages'] if p < 700] or info['pages']
                    return {
                        'code': zero_code,
                        'match_type': f'Khớp nhóm {zero_code}',
                        'status': 'Đã khớp TT38',
                        'appendix': info['appendix'],
                        'pdf_file': info['pdf_file'],
                        'page': pages[0]
                    }
        return None

    def audit_file(self, input_path, output_dir):
        """
        Thẩm tra toàn diện 1 file dự toán:
        1. Phân tích sheet Công trình, Giá tổng hợp, Tổng hợp VT, TH chi phí TB
        2. Tạo sheet Dashboard Tổng hợp thẩm tra
        3. Ghi file Excel mới và trả về thống kê giảm trừ
        """
        os.makedirs(output_dir, exist_ok=True)
        fname = os.path.basename(input_path)
        base_name, _ = os.path.splitext(fname)
        out_fname = f"{base_name}_Audited.xlsx"
        out_fpath = os.path.join(output_dir, out_fname)

        cal_wb = CalamineWorkbook.from_path(input_path)
        sheet_names = cal_wb.sheet_names

        out_wb = openpyxl.Workbook()
        if "Sheet" in out_wb.sheetnames:
            out_wb.remove(out_wb["Sheet"])

        deduction_items = []
        tt38_matched_count = 0
        tt38_tt_count = 0
        tt38_unmatched_count = 0
        total_original_cost = 0.0
        total_deduction_amount = 0.0
        audited_norm_rows = []

        # Xử lý từng sheet
        for sname in sheet_names:
            rows = cal_wb.get_sheet_by_name(sname).to_python()
            out_ws = out_wb.create_sheet(title=sname[:31])

            if not rows:
                continue

            # Sao chép dữ liệu cơ bản
            for r_idx, row in enumerate(rows, 1):
                for c_idx, val in enumerate(row, 1):
                    cell = out_ws.cell(row=r_idx, column=c_idx, value=val)
                    cell.font = FONT_REGULAR
                    cell.border = BORDER_THIN
                    if r_idx <= 4:
                        cell.font = FONT_BOLD

            # 1. Thẩm tra Sheet 'TH chi phí TB' (Thiết bị)
            if sname == 'TH chi phí TB':
                self._audit_sheet_thiet_bi(out_ws, rows, deduction_items)

            # 2. Thẩm tra Sheet 'Tổng hợp VT' (Vật tư chênh lệch giá)
            elif sname == 'Tổng hợp VT':
                self._audit_sheet_tong_hop_vt(out_ws, rows, deduction_items)

            # 3. Thẩm tra Sheet 'Giá tổng hợp'
            elif sname == 'Giá tổng hợp':
                self._audit_sheet_gia_tong_hop(out_ws, rows, deduction_items, audited_norm_rows)

            # 4. Thẩm tra Sheet 'Công trình' (Tra cứu mã TT38)
            elif sname == 'Công trình':
                self._audit_sheet_cong_trinh(out_ws, rows, audited_norm_rows)

            # Lấy tổng giá trị dự toán từ các sheet tổng hợp
            if sname in ['TH chi phí XD', 'THKP hạng mục', 'TH chi phí TB', 'Bang TMDT', 'Phuluc THKP', 'BẢNG DỰ TOÁN']:
                for r in rows:
                    row_str = " ".join(str(c or '') for c in r).lower()
                    for v in r:
                        if isinstance(v, (int, float)) and v > 100000:
                            if any(k in row_str for k in ['tổng cộng (làm tròn)', 'tổng cộng', 'cộng giá trị dự toán', 'gxd', 'tổng chi phí']):
                                total_original_cost = max(total_original_cost, float(v))

        # Đếm các định mức
        for item in audited_norm_rows:
            if item.get('status') == 'Đã khớp TT38':
                tt38_matched_count += 1
            elif item.get('status') == 'Tạm tính':
                tt38_tt_count += 1
            else:
                tt38_unmatched_count += 1

        total_deduction_amount = sum(d['deduction_amount'] for d in deduction_items)
        if total_original_cost == 0.0 and deduction_items:
            total_original_cost = total_deduction_amount * 3  # Fallback ước lượng

        total_audited_cost = max(0.0, total_original_cost - total_deduction_amount)
        reduction_pct = (total_deduction_amount / total_original_cost * 100) if total_original_cost > 0 else 0.0

        # Tạo Sheet 1: 'TỔNG HỢP THẨM TRA' (Dashboard)
        ws_dash = out_wb.create_sheet(title="TỔNG HỢP THẨM TRA", index=0)
        self._build_dashboard_sheet(
            ws_dash, fname, total_original_cost, total_audited_cost,
            total_deduction_amount, reduction_pct,
            deduction_items, audited_norm_rows
        )

        # Định dạng độ rộng cột
        for ws in out_wb.worksheets:
            for col in ws.columns:
                max_len = 0
                for cell in col[:30]:
                    val = str(cell.value or '')
                    if '\n' in val:
                        val = val.split('\n')[0]
                    max_len = max(max_len, len(val))
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

        out_wb.active = 0
        out_wb.save(out_fpath)

        return {
            'file_name': fname,
            'audited_file_path': out_fpath,
            'total_cost_original': total_original_cost,
            'total_cost_audited': total_audited_cost,
            'total_deduction': total_deduction_amount,
            'reduction_pct': reduction_pct,
            'deduction_items': deduction_items,
            'tt38_matched_count': tt38_matched_count,
            'tt38_tt_count': tt38_tt_count,
            'tt38_unmatched_count': tt38_unmatched_count,
            'audited_norm_rows': audited_norm_rows
        }

    def _audit_sheet_thiet_bi(self, ws, rows, deduction_items):
        """Thẩm tra chi phí thiết bị trong Sheet TH chi phí TB"""
        col_name = 1
        col_unit = 2
        col_qty = 3
        col_price = 4
        col_total = 5

        start_col = len(rows[4]) + 1 if len(rows) > 4 else 8
        headers_audit = ["LINK BÁO GIÁ ĐỐI CHIẾU", "ĐƠN GIÁ BÁO GIÁ", "CHÊNH LỆCH ĐƠN GIÁ", "TIỀN GIẢM TRỪ", "KẾT QUẢ THẨM TRA", "GHI CHÚ THẨM TRA"]
        for i, h in enumerate(headers_audit):
            c = ws.cell(row=5, column=start_col + i, value=h)
            c.fill = FILL_HEADER
            c.font = FONT_HEADER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for r_idx in range(7, len(rows) + 1):
            name_val = ws.cell(row=r_idx, column=2).value
            unit_val = ws.cell(row=r_idx, column=3).value
            qty_val = ws.cell(row=r_idx, column=4).value
            price_val = ws.cell(row=r_idx, column=5).value

            if not name_val or not price_val or str(name_val).strip().startswith('*') or 'chi phí' in str(name_val).lower():
                continue

            try:
                qty = float(qty_val or 0)
                price = float(price_val or 0)
            except (ValueError, TypeError):
                continue

            matched_quote, score = self.quotation_indexer.match(str(name_val), str(unit_val or ''))
            if matched_quote:
                quote_price = matched_quote['price']
                diff = price - quote_price
                deduct = diff * qty if diff > 0 else 0.0

                c_link = ws.cell(row=r_idx, column=start_col)
                c_link.value = f"Xem Báo giá: {matched_quote['file_name']}"
                c_link.hyperlink = f"file:///{matched_quote['file_path'].replace('\\', '/')}"
                c_link.font = FONT_LINK

                c_qprice = ws.cell(row=r_idx, column=start_col + 1, value=quote_price)
                c_qprice.number_format = "#,##0"

                c_diff = ws.cell(row=r_idx, column=start_col + 2, value=diff)
                c_diff.number_format = "#,##0"

                c_deduct = ws.cell(row=r_idx, column=start_col + 3, value=deduct)
                c_deduct.number_format = "#,##0"

                c_status = ws.cell(row=r_idx, column=start_col + 4)
                c_note = ws.cell(row=r_idx, column=start_col + 5)

                if diff > 0:
                    c_status.value = "🔴 VƯỢT BÁO GIÁ"
                    c_status.font = FONT_ALERT
                    c_note.value = f"Đơn giá dự toán cao hơn báo giá Daikin {diff:,.0f} đ/cái. Giảm trừ: {deduct:,.0f} đ"
                    for c_idx in range(1, start_col + 6):
                        ws.cell(row=r_idx, column=c_idx).fill = FILL_RED

                    deduction_items.append({
                        'sheet': 'TH chi phí TB',
                        'item_name': str(name_val),
                        'unit': str(unit_val or ''),
                        'qty': qty,
                        'est_price': price,
                        'quote_price': quote_price,
                        'diff_price': diff,
                        'deduction_amount': deduct,
                        'quote_file': matched_quote['file_name'],
                        'quote_path': matched_quote['file_path'],
                        'reason': f"Đơn giá thiết bị cao hơn Báo giá chính hãng Daikin {diff:,.0f} đ/{unit_val}"
                    })
                else:
                    c_status.value = "🟢 PHÙ HỢP BÁO GIÁ"
                    c_status.font = FONT_OK
                    c_note.value = f"Đơn giá phù hợp báo giá Daikin ({quote_price:,.0f} đ)"
                    for c_idx in range(1, start_col + 6):
                        ws.cell(row=r_idx, column=c_idx).fill = FILL_GREEN

    def _audit_sheet_tong_hop_vt(self, ws, rows, deduction_items):
        """Thẩm tra đơn giá vật tư trong Sheet Tổng hợp VT"""
        col_name = 3
        col_unit = 4
        col_qty = 5
        col_price = 8 if len(rows[4]) >= 8 else 6

        start_col = len(rows[4]) + 1 if len(rows) > 4 else 9
        headers_audit = ["LINK BÁO GIÁ ĐỐI CHIẾU", "ĐƠN GIÁ BÁO GIÁ", "CHÊNH LỆCH ĐƠN GIÁ", "TIỀN GIẢM TRỪ", "KẾT QUẢ THẨM TRA", "GHI CHÚ THẨM TRA"]
        for i, h in enumerate(headers_audit):
            c = ws.cell(row=4, column=start_col + i, value=h)
            c.fill = FILL_HEADER
            c.font = FONT_HEADER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for r_idx in range(5, len(rows) + 1):
            name_val = ws.cell(row=r_idx, column=col_name).value
            unit_val = ws.cell(row=r_idx, column=col_unit).value
            qty_val = ws.cell(row=r_idx, column=col_qty).value
            price_val = ws.cell(row=r_idx, column=col_price).value

            if not name_val or not price_val:
                continue

            try:
                qty = float(qty_val or 0)
                price = float(price_val or 0)
            except (ValueError, TypeError):
                continue

            if price <= 0:
                continue

            matched_quote, score = self.quotation_indexer.match(str(name_val), str(unit_val or ''))
            if matched_quote:
                quote_price = matched_quote['price']
                diff = price - quote_price
                deduct = diff * qty if diff > 0 else 0.0

                c_link = ws.cell(row=r_idx, column=start_col)
                c_link.value = f"Báo giá: {matched_quote['file_name']}"
                c_link.hyperlink = f"file:///{matched_quote['file_path'].replace('\\', '/')}"
                c_link.font = FONT_LINK

                c_qprice = ws.cell(row=r_idx, column=start_col + 1, value=quote_price)
                c_qprice.number_format = "#,##0"

                c_diff = ws.cell(row=r_idx, column=start_col + 2, value=diff)
                c_diff.number_format = "#,##0"

                c_deduct = ws.cell(row=r_idx, column=start_col + 3, value=deduct)
                c_deduct.number_format = "#,##0"

                c_status = ws.cell(row=r_idx, column=start_col + 4)
                c_note = ws.cell(row=r_idx, column=start_col + 5)

                if diff > 0:
                    c_status.value = "🔴 VƯỢT BÁO GIÁ"
                    c_status.font = FONT_ALERT
                    c_note.value = f"Đơn giá vật tư cao hơn báo giá {diff:,.0f} đ/{unit_val}. Giảm trừ: {deduct:,.0f} đ"
                    for c_idx in range(1, start_col + 6):
                        ws.cell(row=r_idx, column=c_idx).fill = FILL_RED

                    # Tránh thêm trùng lặp nếu đã thêm ở sheet khác
                    if not any(d['item_name'] == str(name_val) for d in deduction_items):
                        deduction_items.append({
                            'sheet': 'Tổng hợp VT',
                            'item_name': str(name_val),
                            'unit': str(unit_val or ''),
                            'qty': qty,
                            'est_price': price,
                            'quote_price': quote_price,
                            'diff_price': diff,
                            'deduction_amount': deduct,
                            'quote_file': matched_quote['file_name'],
                            'quote_path': matched_quote['file_path'],
                            'reason': f"Đơn giá vật tư cao hơn Báo giá thị trường {diff:,.0f} đ/{unit_val}"
                        })
                else:
                    c_status.value = "🟢 PHÙ HỢP BÁO GIÁ"
                    c_status.font = FONT_OK
                    c_note.value = f"Đơn giá phù hợp báo giá ({quote_price:,.0f} đ)"
                    for c_idx in range(1, start_col + 6):
                        ws.cell(row=r_idx, column=c_idx).fill = FILL_GREEN

    def _audit_sheet_gia_tong_hop(self, ws, rows, deduction_items, audited_norm_rows):
        """Thẩm tra Sheet Giá tổng hợp: kiểm tra định mức và đối chiếu giá vật liệu Kingpost/Mặt dựng"""
        col_code = 2
        col_name = 3
        col_unit = 4
        col_qty = 5
        col_vl = 6

        start_col = len(rows[4]) + 1 if len(rows) > 4 else 8
        headers_audit = ["LINK TRA CỨU TT 38", "LINK BÁO GIÁ", "KẾT QUẢ THẨM TRA", "TIỀN GIẢM TRỪ", "GHI CHÚ"]
        for i, h in enumerate(headers_audit):
            c = ws.cell(row=4, column=start_col + i, value=h)
            c.fill = FILL_HEADER
            c.font = FONT_HEADER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for r_idx in range(5, len(rows) + 1):
            code_val = ws.cell(row=r_idx, column=col_code).value
            name_val = ws.cell(row=r_idx, column=col_name).value
            unit_val = ws.cell(row=r_idx, column=col_unit).value
            qty_val = ws.cell(row=r_idx, column=col_qty).value
            vl_val = ws.cell(row=r_idx, column=col_vl).value

            if not name_val or str(name_val).strip() in ['', 'TỔNG CỘNG']:
                continue

            # 1. Tra cứu TT38
            tt38_info = self.find_tt38_norm(code_val)
            c_tt38 = ws.cell(row=r_idx, column=start_col)
            if tt38_info and tt38_info.get('page'):
                pl_num = get_pl_number(tt38_info.get('appendix'))
                page = tt38_info['page']
                c_tt38.value = f"Mở TT38: Trang {page} ({tt38_info['appendix']})"
                c_tt38.hyperlink = f"tt38://pl{pl_num}/{page}"
                c_tt38.font = FONT_LINK

            # 2. Đối chiếu báo giá vật liệu
            matched_quote, score = self.quotation_indexer.match(str(name_val), str(unit_val or ''))
            c_quote = ws.cell(row=r_idx, column=start_col + 1)
            c_status = ws.cell(row=r_idx, column=start_col + 2)
            c_deduct = ws.cell(row=r_idx, column=start_col + 3)
            c_note = ws.cell(row=r_idx, column=start_col + 4)

            try:
                qty = float(qty_val or 0)
                vl_price = float(vl_val or 0)
            except (ValueError, TypeError):
                qty, vl_price = 0.0, 0.0

            if matched_quote and vl_price > 0:
                quote_price = matched_quote['price']
                diff = vl_price - quote_price
                c_quote.value = f"Báo giá: {matched_quote['file_name']}"
                c_quote.hyperlink = f"file:///{matched_quote['file_path'].replace('\\', '/')}"
                c_quote.font = FONT_LINK

                if diff > 0:
                    deduct = diff * qty
                    c_deduct.value = deduct
                    c_deduct.number_format = "#,##0"
                    c_status.value = "🔴 VƯỢT BÁO GIÁ"
                    c_status.font = FONT_ALERT
                    c_note.value = f"Đơn giá VL cao hơn báo giá {diff:,.0f} đ/{unit_val}. Giảm trừ: {deduct:,.0f} đ"
                    for c_idx in range(1, start_col + 5):
                        ws.cell(row=r_idx, column=c_idx).fill = FILL_RED

                    if not any(d['item_name'] == str(name_val) for d in deduction_items):
                        deduction_items.append({
                            'sheet': 'Giá tổng hợp',
                            'item_name': str(name_val),
                            'unit': str(unit_val or ''),
                            'qty': qty,
                            'est_price': vl_price,
                            'quote_price': quote_price,
                            'diff_price': diff,
                            'deduction_amount': deduct,
                            'quote_file': matched_quote['file_name'],
                            'quote_path': matched_quote['file_path'],
                            'reason': f"Đơn giá vật liệu cao hơn Công bố giá/Báo giá {diff:,.0f} đ/{unit_val}"
                        })
                else:
                    c_status.value = "🟢 PHÙ HỢP TT38 & BÁO GIÁ"
                    c_status.font = FONT_OK
                    for c_idx in range(1, start_col + 5):
                        ws.cell(row=r_idx, column=c_idx).fill = FILL_GREEN
            elif tt38_info and tt38_info.get('status') == 'Đã khớp TT38':
                c_status.value = "🟢 KHỚP TT38"
                c_status.font = FONT_OK
                for c_idx in range(1, start_col + 5):
                    ws.cell(row=r_idx, column=c_idx).fill = FILL_GREEN
            elif tt38_info and tt38_info.get('status') == 'Tạm tính':
                c_status.value = "🟡 MÃ TẠM TÍNH (TT)"
                c_status.font = FONT_WARN
                c_note.value = "Cần phê duyệt phương án thi công và 03 báo giá cạnh tranh"
                for c_idx in range(1, start_col + 5):
                    ws.cell(row=r_idx, column=c_idx).fill = FILL_YELLOW

    def _audit_sheet_cong_trinh(self, ws, rows, audited_norm_rows):
        """Thẩm tra và gắn link TT38 trực tiếp vào Sheet Công trình"""
        col_code = 3
        col_name = 6
        col_unit = 7

        start_col = len(rows[3]) + 1 if len(rows) > 3 else 9
        c_h1 = ws.cell(row=4, column=start_col, value="LINK TRA CỨU TT 38/2026/TT-BXD")
        c_h1.fill = FILL_HEADER
        c_h1.font = FONT_HEADER
        c_h1.alignment = Alignment(horizontal="center", vertical="center")

        c_h2 = ws.cell(row=4, column=start_col + 1, value="ĐÁNH GIÁ ĐỊNH MỨC")
        c_h2.fill = FILL_HEADER
        c_h2.font = FONT_HEADER
        c_h2.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx in range(5, len(rows) + 1):
            code_val = ws.cell(row=r_idx, column=col_code).value
            name_val = ws.cell(row=r_idx, column=col_name).value
            unit_val = ws.cell(row=r_idx, column=col_unit).value

            if not code_val:
                continue

            norm_info = self.find_tt38_norm(code_val)
            if not norm_info:
                continue

            c_link = ws.cell(row=r_idx, column=start_col)
            c_eval = ws.cell(row=r_idx, column=start_col + 1)

            if norm_info.get('page'):
                pl_num = get_pl_number(norm_info.get('appendix'))
                page = norm_info['page']
                c_link.value = f"Mở TT38: {norm_info['appendix']} - Trang {page}"
                c_link.hyperlink = f"tt38://pl{pl_num}/{page}"
                c_link.font = FONT_LINK
                c_eval.value = f"🟢 Chuẩn TT38 ({norm_info['match_type']})"
                c_eval.font = FONT_OK
            elif norm_info.get('status') == 'Tạm tính':
                c_link.value = "Công tác Tạm tính (Không có trong TT38)"
                c_link.font = FONT_REGULAR
                c_eval.value = "🟡 Tạm tính (Yêu cầu 3 báo giá cạnh tranh)"
                c_eval.font = FONT_WARN
            else:
                c_link.value = "Chưa tìm thấy mã định mức"
                c_eval.value = "🔴 Cần chuẩn hóa định mức"
                c_eval.font = FONT_ALERT

            audited_norm_rows.append({
                'code': str(code_val),
                'name': str(name_val or ''),
                'unit': str(unit_val or ''),
                'status': norm_info.get('status'),
                'appendix': norm_info.get('appendix'),
                'page': norm_info.get('page')
            })

    def _build_dashboard_sheet(self, ws, fname, total_orig, total_audit, total_deduct, pct_deduct, deduction_items, norm_items):
        """Xây dựng Dashboard Tổng hợp thẩm tra chuyên nghiệp"""
        ws.merge_cells("A1:J1")
        ws["A1"] = f"BẢNG TỔNG HỢP KẾT QUẢ THẨM TRA DỰ TOÁN - ĐỐI CHIẾU THÔNG TƯ 38/2026/TT-BXD & BÁO GIÁ"
        ws["A1"].font = FONT_TITLE
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        ws.merge_cells("A2:J2")
        ws["A2"] = f"Hồ sơ dự toán: {fname} | Căn cứ: Luật Xây dựng, TT 38/2026/TT-BXD, Báo giá nhà sản xuất & Công bố giá địa phương"
        ws["A2"].font = Font(name="Times New Roman", size=11, italic=True)
        ws["A2"].alignment = Alignment(horizontal="center")

        # KPI Cards (Dòng 4-5)
        kpis = [
            ("A4:B4", "A5:B5", "TỔNG DỰ TOÁN TRƯỚC THẨM TRA", f"{total_orig:,.0f} đ", "2F5597"),
            ("C4:E4", "C5:E5", "GIÁ TRỊ DỰ TOÁN SAU THẨM TRA", f"{total_audit:,.0f} đ", "375623"),
            ("F4:H4", "F5:H5", "GIÁ TRỊ GIẢM TRỪ THẨM TRA", f"{total_deduct:,.0f} đ", "C00000"),
            ("I4:J4", "I5:J5", "TỶ LỆ GIẢM TRỪ CHI PHÍ", f"{pct_deduct:.2f} %", "B25900")
        ]
        for m1, m2, title, val, color_hex in kpis:
            ws.merge_cells(m1)
            c1 = ws[m1.split(':')[0]]
            c1.value = title
            c1.fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
            c1.font = FONT_HEADER
            c1.alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells(m2)
            c2 = ws[m2.split(':')[0]]
            c2.value = val
            c2.font = Font(name="Times New Roman", size=13, bold=True, color=color_hex)
            c2.alignment = Alignment(horizontal="center", vertical="center")
            c2.border = BORDER_THIN

        ws.row_dimensions[4].height = 24
        ws.row_dimensions[5].height = 28

        # Phần 1: BẢNG TỔNG HỢP CHI TIẾT CÁC KHOẢN GIẢM TRỪ CHI PHÍ
        ws.merge_cells("A7:J7")
        ws["A7"] = "I. CHI TIẾT CÁC HẠNG MỤC SAI LỆCH VÀ KIẾN NGHỊ GIẢM TRỪ CHI PHÍ (MỤC 5.3)"
        ws["A7"].font = FONT_BOLD
        ws["A7"].fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

        deduct_headers = [
            "STT", "Tên hạng mục / Vật tư / Thiết bị", "ĐVT", "Khối lượng",
            "Đơn giá Dự toán (đ)", "Đơn giá Báo giá (đ)", "Chênh lệch (đ)",
            "Thành tiền Giảm trừ (đ)", "Link Báo giá đối chiếu", "Lý do và Căn cứ xử lý"
        ]
        for c_idx, h in enumerate(deduct_headers, 1):
            cell = ws.cell(row=8, column=c_idx, value=h)
            cell.fill = FILL_HEADER
            cell.font = FONT_HEADER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[8].height = 26

        curr_row = 9
        if not deduction_items:
            ws.merge_cells(f"A{curr_row}:J{curr_row}")
            ws[f"A{curr_row}"] = "✓ Không phát hiện sai lệch đơn giá hoặc vượt báo giá trong hồ sơ này."
            ws[f"A{curr_row}"].font = FONT_OK
            ws[f"A{curr_row}"].alignment = Alignment(horizontal="center")
            curr_row += 1
        else:
            for idx, d in enumerate(deduction_items, 1):
                ws.cell(row=curr_row, column=1, value=idx).alignment = Alignment(horizontal="center")
                ws.cell(row=curr_row, column=2, value=d['item_name'])
                ws.cell(row=curr_row, column=3, value=d['unit']).alignment = Alignment(horizontal="center")

                c_qty = ws.cell(row=curr_row, column=4, value=d['qty'])
                c_qty.number_format = "#,##0.00"

                c_ep = ws.cell(row=curr_row, column=5, value=d['est_price'])
                c_ep.number_format = "#,##0"

                c_qp = ws.cell(row=curr_row, column=6, value=d['quote_price'])
                c_qp.number_format = "#,##0"

                c_dp = ws.cell(row=curr_row, column=7, value=d['diff_price'])
                c_dp.number_format = "#,##0"
                c_dp.font = FONT_ALERT

                c_da = ws.cell(row=curr_row, column=8, value=d['deduction_amount'])
                c_da.number_format = "#,##0"
                c_da.font = FONT_ALERT

                c_link = ws.cell(row=curr_row, column=9, value=f"Mở Báo giá: {d['quote_file']}")
                c_link.hyperlink = f"file:///{d['quote_path'].replace('\\', '/')}"
                c_link.font = FONT_LINK

                ws.cell(row=curr_row, column=10, value=d['reason'])

                for c_i in range(1, 11):
                    ws.cell(row=curr_row, column=c_i).border = BORDER_THIN
                    ws.cell(row=curr_row, column=c_i).fill = FILL_RED

                curr_row += 1

            # Dòng tổng cộng giảm trừ
            ws.merge_cells(f"A{curr_row}:G{curr_row}")
            c_sum_lbl = ws.cell(row=curr_row, column=1, value="TỔNG CỘNG GIÁ TRỊ GIẢM TRỪ THẨM TRA:")
            c_sum_lbl.font = FONT_BOLD
            c_sum_lbl.alignment = Alignment(horizontal="right")
            c_sum_lbl.fill = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid")

            c_sum_val = ws.cell(row=curr_row, column=8, value=total_deduct)
            c_sum_val.font = Font(name="Times New Roman", size=11, bold=True, color="C00000")
            c_sum_val.number_format = "#,##0"
            c_sum_val.fill = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid")

            ws.cell(row=curr_row, column=9).fill = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid")
            ws.cell(row=curr_row, column=10).fill = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid")
            curr_row += 2

        # Phần 2: DANH MỤC TRA CỨU ĐỊNH MỨC TT 38/2026/TT-BXD
        ws.merge_cells(f"A{curr_row}:J{curr_row}")
        ws[f"A{curr_row}"] = "II. TỔNG HỢP ĐỐI CHIẾU ĐỊNH MỨC XÂY DỰNG THEO THÔNG TƯ 38/2026/TT-BXD (MỤC 5.2)"
        ws[f"A{curr_row}"].font = FONT_BOLD
        ws[f"A{curr_row}"].fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        curr_row += 1

        norm_headers = [
            "STT", "Mã hiệu dự toán", "Tên công tác xây dựng", "ĐVT",
            "Phụ lục TT38", "Trang PDF", "Đường link Tra cứu TT38 (Nhảy đúng trang)",
            "Đánh giá tính hợp lý", "Tình trạng", "Ghi chú kiểm tra"
        ]
        for c_idx, h in enumerate(norm_headers, 1):
            cell = ws.cell(row=curr_row, column=c_idx, value=h)
            cell.fill = FILL_SUBHEADER
            cell.font = FONT_HEADER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[curr_row].height = 26
        curr_row += 1

        seen_codes = set()
        stt_norm = 1
        for item in norm_items:
            code = item['code']
            if code in seen_codes:
                continue
            seen_codes.add(code)

            ws.cell(row=curr_row, column=1, value=stt_norm).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=2, value=code).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=3, value=item['name'])
            ws.cell(row=curr_row, column=4, value=item['unit']).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=5, value=item.get('appendix') or 'N/A').alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=6, value=item.get('page') or 'N/A').alignment = Alignment(horizontal="center")

            c_link = ws.cell(row=curr_row, column=7)
            c_eval = ws.cell(row=curr_row, column=8)
            c_stat = ws.cell(row=curr_row, column=9)
            c_comm = ws.cell(row=curr_row, column=10)

            if item.get('page'):
                pl_num = get_pl_number(item.get('appendix'))
                page = item['page']
                c_link.value = f"Mở TT38: Trang {page}"
                c_link.hyperlink = f"tt38://pl{pl_num}/{page}"
                c_link.font = FONT_LINK
                c_eval.value = "Áp dụng đúng định mức quy chuẩn"
                c_stat.value = "🟢 Phù hợp"
                c_stat.font = FONT_OK
                c_comm.value = f"Khớp phụ lục {item.get('appendix')}"
                row_fill = FILL_GREEN
            elif item.get('status') == 'Tạm tính':
                c_link.value = "Mã TT (Không có trong TT38)"
                c_eval.value = "Công tác tạm tính"
                c_stat.value = "🟡 Cần rà soát"
                c_stat.font = FONT_WARN
                c_comm.value = "Cần phê duyệt biện pháp & 03 báo giá"
                row_fill = FILL_YELLOW
            else:
                c_link.value = "Chưa có link"
                c_eval.value = "Mã hiệu chưa khớp TT38"
                c_stat.value = "🔴 Cần chuẩn hóa"
                c_stat.font = FONT_ALERT
                c_comm.value = "Kiểm tra lại danh mục TT38"
                row_fill = FILL_RED

            for c_i in range(1, 11):
                ws.cell(row=curr_row, column=c_i).border = BORDER_THIN
                ws.cell(row=curr_row, column=c_i).fill = row_fill

            curr_row += 1
            stt_norm += 1
