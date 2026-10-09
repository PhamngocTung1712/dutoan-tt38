import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

folder = r"d:\Cong ty\DuToanChuyenNghiep\Bao_Gia_Dau_Vao"
os.makedirs(folder, exist_ok=True)

# 1. Báo giá Vách kính và Mặt dựng
wb1 = openpyxl.Workbook()
ws1 = wb1.active
ws1.title = "Bao_Gia_Mat_Dung"

header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
regular_font = Font(name="Times New Roman", size=11)
bold_font = Font(name="Times New Roman", size=11, bold=True)

ws1.merge_cells("A1:G1")
ws1["A1"] = "CÔNG TY CP NHÔM KÍNH & MẶT DỰNG HÀ NỘI - BẢNG BÁO GIÁ VẬT TƯ THI CÔNG"
ws1["A1"].font = Font(name="Times New Roman", size=14, bold=True, color="1F4E79")
ws1["A1"].alignment = Alignment(horizontal="center")

ws1.merge_cells("A2:G2")
ws1["A2"] = "Số: 128/2026/BG-MDHN | Ngày hiệu lực: 15/08/2026 | Dự án: Nhà ở CBCS Công an - NO1 Cầu Giấy"
ws1["A2"].font = Font(name="Times New Roman", size=10, italic=True)
ws1["A2"].alignment = Alignment(horizontal="center")

headers = ["STT", "Mã vật tư", "Tên vật tư / Quy cách kỹ thuật", "Đơn vị tính", "Đơn giá trước thuế (VNĐ)", "Nhà sản xuất / Xuất xứ", "Ghi chú"]
for c, h in enumerate(headers, 1):
    cell = ws1.cell(row=4, column=c, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

data_matdung = [
    [1, "VT.MD01", "Vách kính khung nhôm mặt tiền (Nhôm Xingfa hệ 65, kính hộp dán an toàn Low-E 24mm)", "m2", 4850000, "Xingfa / Kính Viglacera", "Báo giá cạnh tranh 1"],
    [2, "VT.MD02", "Tấm ốp mặt tiền (Tấm hợp kim nhôm Aluminium composite Alcorest chống cháy dày 4mm)", "m2", 3650000, "Alcorest Việt Dũng", "Báo giá cạnh tranh 1"],
    [3, "VT.MD03", "Keo silicone kết cấu kết dính mặt dựng Dowsil 791 / 995", "Tuýp", 145000, "Dow Corning (Mỹ)", "Chống thấm kết cấu"],
    [4, "VT.MD04", "Bản mã thép mạ kẽm nhúng nóng liên kết hệ mặt dựng", "kg", 28500, "Thép Hòa Phát", "Đã bao gồm mạ kẽm"],
    [5, "VT.MD05", "Bulong nở hóa chất Fischer M16x190", "Bộ", 85000, "Fischer (Đức)", "Kèm ống nhộng"]
]

for r_idx, row in enumerate(data_matdung, 5):
    for c_idx, val in enumerate(row, 1):
        cell = ws1.cell(row=r_idx, column=c_idx, value=val)
        cell.font = regular_font
        if c_idx == 5:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right")
        elif c_idx in [1, 2, 4]:
            cell.alignment = Alignment(horizontal="center")

for col in ws1.columns:
    max_len = max(len(str(c.value or '')) for c in col)
    ws1.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = max(max_len + 4, 12)

wb1.save(os.path.join(folder, "01_Bao_Gia_Mat_Dung_Nhom_Kinh_2026.xlsx"))

# 2. Báo giá Hệ thống Điều hòa không khí VRF
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = "Bao_Gia_Dieu_Hoa_VRF"

ws2.merge_cells("A1:G1")
ws2["A1"] = "CÔNG TY TNHH KỸ THUẬT LẠNH & ĐIỀU HÒA DAIKIN MIỀN BẮC - BẢNG GIÁ THIẾT BỊ VRV/VRF"
ws2["A1"].font = Font(name="Times New Roman", size=14, bold=True, color="1F4E79")
ws2["A1"].alignment = Alignment(horizontal="center")

ws2.merge_cells("A2:G2")
ws2["A2"] = "Số: 452/DK-2026 | Ngày phát hành: 10/09/2026 | Áp dụng cho khối đế NO1 Cầu Giấy"
ws2["A2"].font = Font(name="Times New Roman", size=10, italic=True)
ws2["A2"].alignment = Alignment(horizontal="center")

headers2 = ["STT", "Mã thiết bị", "Tên thiết bị / Chủng loại công suất", "Đơn vị tính", "Đơn giá trước thuế (VNĐ)", "Hãng sản xuất", "Bảo hành"]
for c, h in enumerate(headers2, 1):
    cell = ws2.cell(row=4, column=c, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

data_vrf = [
    [1, "TB.VRF10", "Dàn nóng điều hòa trung tâm VRF công suất 10 HP", "Cái", 148000000, "Daikin / Thái Lan", "24 tháng"],
    [2, "TB.VRF14", "Dàn nóng điều hòa trung tâm VRF công suất 14 HP", "Cái", 205000000, "Daikin / Thái Lan", "24 tháng"],
    [3, "TB.VRF16", "Dàn nóng điều hòa trung tâm VRF công suất 16 HP", "Cái", 235000000, "Daikin / Thái Lan", "24 tháng"],
    [4, "TB.VRF26", "Dàn nóng điều hòa trung tâm VRF công suất 26 HP", "Cái", 388000000, "Daikin / Thái Lan", "24 tháng"],
    [5, "TB.VRF40", "Dàn nóng điều hòa trung tâm VRF công suất 40 HP", "Cái", 565000000, "Daikin / Thái Lan", "24 tháng"],
    [6, "TB.VRF46", "Dàn nóng điều hòa trung tâm VRF công suất 46 HP", "Cái", 685000000, "Daikin / Thái Lan", "24 tháng"],
    [7, "TB.DL09", "Dàn lạnh cassette âm trần 4 hướng thổi CSL: 9.0 kW", "Cái", 12500000, "Daikin / Thái Lan", "24 tháng"],
    [8, "TB.DL11", "Dàn lạnh cassette âm trần 4 hướng thổi CSL: 11.2 kW", "Cái", 14200000, "Daikin / Thái Lan", "24 tháng"],
    [9, "TB.DL14", "Dàn lạnh cassette âm trần 4 hướng thổi CSL: 14.0 kW", "Cái", 16800000, "Daikin / Thái Lan", "24 tháng"],
    [10, "TB.DLO25", "Dàn lạnh giấu trần nối ống gió CSL: 2.5 kW", "Cái", 8900000, "Daikin / Thái Lan", "24 tháng"],
    [11, "TB.DLO71", "Dàn lạnh giấu trần nối ống gió CSL: 7.1 kW", "Cái", 13500000, "Daikin / Thái Lan", "24 tháng"]
]

for r_idx, row in enumerate(data_vrf, 5):
    for c_idx, val in enumerate(row, 1):
        cell = ws2.cell(row=r_idx, column=c_idx, value=val)
        cell.font = regular_font
        if c_idx == 5:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right")
        elif c_idx in [1, 2, 4]:
            cell.alignment = Alignment(horizontal="center")

for col in ws2.columns:
    max_len = max(len(str(c.value or '')) for c in col)
    ws2.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = max(max_len + 4, 12)

wb2.save(os.path.join(folder, "02_Bao_Gia_Thiet_Bi_Dieu_Hoa_VRF_2026.xlsx"))

# 3. Báo giá Thép hình Kingpost & Vật liệu xây dựng phần ngầm
wb3 = openpyxl.Workbook()
ws3 = wb3.active
ws3.title = "Bao_Gia_Vat_Lieu_Ngam"

ws3.merge_cells("A1:G1")
ws3["A1"] = "TỔNG CÔNG TY KINH DOANH THÉP & KIM KHÍ - BẢNG GIÁ THÉP HÌNH & VẬT LIỆU"
ws3["A1"].font = Font(name="Times New Roman", size=14, bold=True, color="1F4E79")
ws3["A1"].alignment = Alignment(horizontal="center")

ws3.merge_cells("A2:G2")
ws3["A2"] = "Số: 89/KKHN-2026 | Công bố giá Quý 3/2026 tại thị trường Hà Nội"
ws3["A2"].font = Font(name="Times New Roman", size=10, italic=True)
ws3["A2"].alignment = Alignment(horizontal="center")

headers3 = ["STT", "Mã vật tư", "Tên vật tư / Quy cách", "Đơn vị tính", "Đơn giá trước thuế (VNĐ)", "Tiêu chuẩn / Nhà SX", "Ghi chú"]
for c, h in enumerate(headers3, 1):
    cell = ws3.cell(row=4, column=c, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

data_ngam = [
    [1, "VT.TH01", "Thép hình H300, H350, H400 gia công cọc Kingpost", "tấn", 18500000, "SS400 / Posco", "Thép Posco Yamato"],
    [2, "VT.DC01", "Đinh chống cắt (Stud bolt) D19, dài L=100mm kèm vòng đệm gốm", "cái", 32000, "Nelson / Nhập khẩu", "Báo giá cạnh tranh"],
    [3, "VT.VC01", "Vữa không co ngót cường độ cao B40 (SikaGrout 214-11)", "m3", 21500000, "Sika Việt Nam", "Đơn giá quy đổi m3 vữa"],
    [4, "VT.NL01", "Nilon lót chống mất nước bê tông lót (bạt nilon tái sinh dày 0.05mm)", "100m2", 680000, "Việt Nam", "Đạt TCVN"],
    [5, "VT.BT150", "Bê tông thương phẩm mác 150 (M150) độ sụt 12±2, đá 1x2", "m3", 1150000, "Bê tông Việt Hàn", "Công bố giá liên sở Q3/2026"]
]

for r_idx, row in enumerate(data_ngam, 5):
    for c_idx, val in enumerate(row, 1):
        cell = ws3.cell(row=r_idx, column=c_idx, value=val)
        cell.font = regular_font
        if c_idx == 5:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right")
        elif c_idx in [1, 2, 4]:
            cell.alignment = Alignment(horizontal="center")

for col in ws3.columns:
    max_len = max(len(str(c.value or '')) for c in col)
    ws3.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = max(max_len + 4, 12)

wb3.save(os.path.join(folder, "03_Cong_Bo_Gia_Thep_Va_Vat_Lieu_Xay_Dung_2026.xlsx"))

print("Đã tạo xong bộ 3 file Báo giá mẫu trong Bao_Gia_Dau_Vao!")
